"""Read Excel/CSV uploads into a list of dict rows with normalised (lowercase) headers."""
import csv
import io

import openpyxl


class UploadError(Exception):
    pass


def _norm(value):
    return "" if value is None else str(value).strip()


def read_rows(uploaded_file):
    """Return (headers, rows). `rows` is a list of dicts keyed by lowercase header."""
    name = uploaded_file.name.lower()
    data = uploaded_file.read()
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        grid = list(csv.reader(io.StringIO(text)))
    elif name.endswith((".xlsx", ".xlsm")):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise UploadError("Could not read that Excel file.") from exc
        grid = [list(r) for r in wb.active.iter_rows(values_only=True)]
    else:
        raise UploadError("Please upload a .csv or .xlsx file.")
    grid = [[_norm(c) for c in row] for row in grid if any(_norm(c) for c in row)]
    if not grid:
        raise UploadError("The file is empty.")
    headers = [h.lower().replace(" ", "_") for h in grid[0]]
    rows = []
    for line_no, row in enumerate(grid[1:], start=2):
        row = row + [""] * (len(headers) - len(row))
        d = dict(zip(headers, row))
        d["_line"] = line_no
        rows.append(d)
    return headers, rows


def template_csv(headers, example_rows=()):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    w.writerows(example_rows)
    return buf.getvalue()


def template_xlsx(headers, example_rows=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in example_rows:
        ws.append(list(r))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
