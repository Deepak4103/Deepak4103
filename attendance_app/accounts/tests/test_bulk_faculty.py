import csv
import io

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from accounts import bulk
from accounts.models import User

HEADER = "full_name,user_id,password\n"


class BulkFacultyTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user("adm", password="pw123456", full_name="Admin", role=User.ADMIN)
        self.client.force_login(self.admin)
        self.url = reverse("faculty_upload")

    def upload(self, text, name="f.csv"):
        return self.client.post(self.url, {"file": SimpleUploadedFile(name, text.encode())})

    def confirm(self):
        return self.client.post(self.url, {"action": "confirm"})

    # ---- preview and confirm
    def test_preview_saves_nothing_then_confirm_creates_accounts(self):
        r = self.upload(HEADER + "Asha Rao,asha,Secret#123\nRavi Kumar,ravi,Another#456\n")
        self.assertContains(r, "2 will be created")
        self.assertEqual(User.objects.filter(role=User.FACULTY).count(), 0)       # preview saves nothing
        r = self.confirm()
        self.assertContains(r, "2 accounts created")
        asha = User.objects.get(username="asha")
        self.assertEqual((asha.full_name, asha.role, asha.is_active), ("Asha Rao", User.FACULTY, True))
        self.assertTrue(asha.check_password("Secret#123"))
        self.assertTrue(asha.must_change_password)

    def test_passwords_are_hashed_not_stored_in_clear(self):
        self.upload(HEADER + "Asha Rao,asha,Secret#123\n")
        self.confirm()
        self.assertNotIn("Secret#123", User.objects.get(username="asha").password)

    def test_blank_password_is_generated_and_shown_once(self):
        self.upload(HEADER + "Asha Rao,asha,\n")
        r = self.confirm()
        pw = r.context["created"][0][2]
        self.assertGreaterEqual(len(pw), 8)
        self.assertContains(r, pw)
        self.assertTrue(User.objects.get(username="asha").check_password(pw))
        # the passwords are not kept anywhere to be shown again
        self.client.force_login(self.admin)
        self.assertNotContains(self.client.get(reverse("faculty_list")), pw)
        again = self.confirm()
        self.assertRedirects(again, self.url)

    def test_generated_passwords_differ_and_use_clear_characters(self):
        a, b = bulk.generate_password(), bulk.generate_password()
        self.assertNotEqual(a, b)
        self.assertFalse(set(a) & set("0O1lI"))

    def test_new_faculty_can_log_in_and_must_change_password(self):
        self.upload(HEADER + "Asha Rao,asha,Secret#123\n")
        self.confirm()
        self.client.logout()
        self.assertTrue(self.client.login(username="asha", password="Secret#123"))
        r = self.client.get(reverse("faculty_home"))
        self.assertRedirects(r, reverse("password_change"))

    def test_download_link_has_a_valid_csv(self):
        self.upload(HEADER + '"Rao, Asha",asha,Secret#123\n')
        r = self.confirm()
        rows = list(csv.reader(io.StringIO(r.context["csv"])))
        self.assertEqual(rows, [["full_name", "user_id", "password"], ["Rao, Asha", "asha", "Secret#123"]])

    # ---- validation
    def test_problem_rows_are_skipped_and_explained(self):
        User.objects.create_user("taken", password="x", full_name="Taken", role=User.FACULTY)
        r = self.upload(HEADER + ("Good One,good,Secret#123\n"
                                  ",noname,Secret#123\n"                  # no name
                                  "No Id,,Secret#123\n"                    # no user id
                                  "Has Space,has space,Secret#123\n"       # space in id
                                  "Bad Chars,bad#id,Secret#123\n"          # illegal character
                                  "Dup One,good,Secret#123\n"              # duplicate in file
                                  "Case Dup,GOOD,Secret#123\n"             # duplicate ignoring case
                                  "Exists,TAKEN,Secret#123\n"              # already an account (any case)
                                  "Admin Clash,adm,Secret#123\n"           # clashes with the admin
                                  "Weak,weak,123\n"))                      # too short
        self.assertContains(r, "1 will be created")
        self.assertContains(r, "9 will be skipped")
        for msg in ("Name and User ID are both required", "no spaces", "only letters", "Duplicate User ID",
                    "already exists", "Password:"):
            self.assertContains(r, msg)
        self.confirm()
        self.assertEqual(sorted(User.objects.filter(role=User.FACULTY).values_list("username", flat=True)),
                         ["good", "taken"])

    def test_duplicate_userid_in_file_keeps_only_the_first(self):
        valid, errors = bulk.validate_rows([
            {"_line": 2, "full_name": "A", "user_id": "x", "password": ""},
            {"_line": 3, "full_name": "B", "user_id": "x", "password": ""}])
        self.assertEqual([v["full_name"] for v in valid], ["A"])
        self.assertEqual(errors[0][0], 3)

    def test_recheck_at_save_time_prevents_duplicates(self):
        self.upload(HEADER + "Asha Rao,asha,Secret#123\n")
        User.objects.create_user("asha", password="x", full_name="Someone else", role=User.FACULTY)   # created meanwhile
        self.confirm()
        self.assertEqual(User.objects.filter(username="asha").count(), 1)
        self.assertEqual(User.objects.get(username="asha").full_name, "Someone else")

    def test_header_aliases_and_case(self):
        r = self.upload("Name,Username,Password\nAsha Rao,asha,Secret#123\n")
        self.assertContains(r, "1 will be created")

    def test_missing_headers_are_rejected(self):
        r = self.upload("a,b\n1,2\n")
        self.assertContains(r, "Header row must contain")

    def test_unsupported_file_type(self):
        r = self.upload("x", name="f.txt")
        self.assertContains(r, "csv or .xlsx")

    def test_excel_file_with_numeric_cells(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["full_name", "user_id", "password"])
        ws.append(["Asha Rao", 1001, 654321])          # numbers typed into Excel cells
        buf = io.BytesIO()
        wb.save(buf)
        r = self.client.post(self.url, {"file": SimpleUploadedFile("f.xlsx", buf.getvalue())})
        self.assertContains(r, "1 will be created")
        self.confirm()
        self.assertTrue(User.objects.get(username="1001").check_password("654321"))

    def test_empty_file_and_blank_lines(self):
        self.assertContains(self.upload(""), "empty")
        r = self.upload(HEADER + "\n\nAsha Rao,asha,Secret#123\n\n")
        self.assertContains(r, "1 will be created")

    def test_nothing_valid_means_no_create_button(self):
        r = self.upload(HEADER + "Bad,bad id,x\n")
        self.assertNotContains(r, "Create 0")

    # ---- templates and access
    def test_blank_templates(self):
        r = self.client.get(reverse("faculty_template", args=["csv"]))
        self.assertEqual(r.content.decode().strip(), "full_name,user_id,password")
        r = self.client.get(reverse("faculty_template", args=["xlsx"]))
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[1]], bulk.HEADERS)
        self.assertEqual(ws.max_row, 1)

    def test_faculty_list_links_to_the_upload(self):
        self.assertContains(self.client.get(reverse("faculty_list")), "Bulk upload")

    def test_only_admin_can_use_it(self):
        fac = User.objects.create_user("fac", password="pw123456", full_name="Fac", role=User.FACULTY)
        self.client.force_login(fac)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertEqual(self.client.post(self.url, {"action": "confirm"}).status_code, 403)
        self.assertEqual(self.client.get(reverse("faculty_template", args=["csv"])).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)

    def test_expired_upload(self):
        r = self.confirm()
        self.assertRedirects(r, self.url)
