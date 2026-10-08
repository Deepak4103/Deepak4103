"""Excel and PDF output for a Report."""
import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
from xml.sax.saxutils import escape


FONT_DIR = Path(__file__).resolve().parent / "fonts"


def _register_fonts():
    """Use the bundled DejaVu Sans (accented Latin, Greek, Cyrillic and more) for PDFs.
    Falls back to the built-in Helvetica if the font files are missing."""
    try:
        if "DejaVuSans" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DejaVuSans", str(FONT_DIR / "DejaVuSans.ttf")))
            pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
        return "DejaVuSans", "DejaVuSans-Bold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


def to_xlsx(report):
    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:31].replace("/", "-")
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([report.subtitle])
    ws.append([f"Generated {datetime.datetime.now():%d-%m-%Y %H:%M}"])
    ws.append([])
    ws.append(list(report.columns))
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2456D6")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(report.rows):
        ws.append(list(row))
        if report.flag(i) == "bad":
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="FDECEB")
    for n in report.notes:
        ws.append([])
        ws.append([n])
    widths = [len(str(c)) for c in report.columns]
    for row in report.rows:
        for j, v in enumerate(row):
            widths[j] = max(widths[j], min(len(str(v)), 45))
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(w + 2, 8), 47)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(report):
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title=report.title)
    styles = getSampleStyleSheet()
    regular, bold = _register_fonts()
    for name in ("Title", "Normal", "BodyText", "Italic"):
        styles[name].fontName = bold if name == "Title" else regular
    n_cols = max(len(report.columns), 1)
    size = 8 if n_cols <= 7 else 7 if n_cols <= 10 else 6
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontName=regular, fontSize=size, leading=size + 2)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName=bold)
    story = [Paragraph(escape(report.title), styles["Title"]),
             Paragraph(escape(report.subtitle), styles["Normal"]),
             Paragraph(f"Generated {datetime.datetime.now():%d-%m-%Y %H:%M}", styles["Normal"]), Spacer(1, 4 * mm)]
    data = [[Paragraph(escape(str(c)), head) for c in report.columns]]
    for row in report.rows:
        data.append([Paragraph(escape(str(v)), cell) for v in row])
    # column widths in proportion to the longest text (bounded), so narrow columns stay narrow
    avail = landscape(A4)[0] - 20 * mm
    def weight(j, col):
        longest_word = max((len(w) for w in str(col).split()), default=4)
        longest_cell = max((len(str(r[j])) for r in report.rows), default=0)
        return max(longest_word + 3, min(longest_cell, 40) + 2, 6)   # never narrower than its header word
    weights = [weight(j, c) for j, c in enumerate(report.columns)]
    total = sum(weights)
    table = LongTable(data, colWidths=[avail * w / total for w in weights], repeatRows=1)
    style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2456D6")),
             ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#bbbbbb")),
             ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f8fb")])]
    for i in range(len(report.rows)):
        if report.flag(i) == "bad":
            style.append(("BACKGROUND", (0, i + 1), (-1, i + 1), colors.HexColor("#FDECEB")))
    table.setStyle(TableStyle(style))
    story.append(table if report.rows else Paragraph("No records for these filters.", styles["Normal"]))
    for n in report.notes:
        story += [Spacer(1, 2 * mm), Paragraph(escape(n), styles["Italic"])]
    doc.build(story)
    return out.getvalue()
