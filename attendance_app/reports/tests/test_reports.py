import datetime
import io

import openpyxl
from django.urls import reverse

from academics.models import Holiday, set_setting
from attendance import services as att
from leaves import services as leave_svc
from leaves.models import Adjustment, LeaveRequest
from reports import builders
from reports.stats import Filters

from .base import ReportTestCase, week


def F(**kw):
    return Filters(**kw)


class StudentWiseTests(ReportTestCase):
    def test_per_subject_and_overall(self):
        self.held(week(0), absent=[self.r(1)])
        self.held(week(1))
        self.held(week(0), period=2, absent=[self.r(1)])
        rep = builders.student_wise(F(school_class=self.c))
        self.assertIn("Data Structures (DS)", rep.columns[2])
        row = next(r for r in rep.rows if r[0] == "R1")
        self.assertEqual(row[2:], ["1/2 (50.0%)", "0/1 (0.0%)", 3, 1, "33.3%"])
        self.assertEqual(next(r for r in rep.rows if r[0] == "R2")[2:], ["2/2 (100.0%)", "1/1 (100.0%)", 3, 3, "100.0%"])

    def test_students_below_threshold_are_flagged(self):
        self.held(week(0), absent=[self.r(1)])
        rep = builders.student_wise(F(school_class=self.c))
        self.assertEqual([rep.flag(i) for i, r in enumerate(rep.rows) if r[0] == "R1"], ["bad"])
        self.assertEqual([rep.flag(i) for i, r in enumerate(rep.rows) if r[0] == "R2"], [""])

    def test_no_classes_held_shows_dashes(self):
        row = builders.student_wise(F(school_class=self.c)).rows[0]
        self.assertEqual(row[2:], ["-", "-", 0, 0, "-"])

    def test_subject_filter_limits_columns(self):
        rep = builders.student_wise(F(school_class=self.c, subject_code="db"))
        self.assertEqual(len(rep.columns), 2 + 1 + 3)


class SummaryTests(ReportTestCase):
    def test_class_and_subject_summary_with_overall_row(self):
        self.held(week(0), absent=[self.r(1)], user=self.ravi)                    # 3 present, 1 absent
        self.held(week(1), user=self.ravi)                                        # 4 present
        self.held(week(0), period=2, absent=[self.r(1), self.r(2)], user=self.sneha)   # 2 present, 2 absent
        rep = builders.class_subject_summary(F())
        ds = next(r for r in rep.rows if "DS" in r[1])
        self.assertEqual(ds, ["A", "Data Structures (DS)", "Ravi", 2, 7, 1, "87.5%"])
        db = next(r for r in rep.rows if "DB" in r[1])
        self.assertEqual(db, ["A", "Databases (DB)", "Sneha", 1, 2, 2, "50.0%"])
        self.assertEqual(rep.rows[-1], ["A", "ALL SUBJECTS", "", 3, 9, 3, "75.0%"])

    def test_filters(self):
        self.held(week(0))
        self.held(week(0), period=2)
        rep = builders.class_subject_summary(F(subject_code="DB"))
        self.assertEqual([r[1] for r in rep.rows], ["Databases (DB)", "ALL SUBJECTS"])


class DailyTests(ReportTestCase):
    def test_entered_and_pending_periods_with_topic(self):
        self.set_today(week(0))
        self.held(week(0), absent=[self.r(1)], topic="Stacks", user=self.ravi)
        rep = builders.daily(F(date=week(0)))
        p1, p2 = rep.rows
        self.assertEqual((p1[1], p1[3], p1[4], p1[5], p1[6], p1[7], p1[8]),
                         (1, "Data Structures", "Ravi", "Entered", 3, 1, "Stacks"))
        self.assertEqual((p2[1], p2[5], p2[6]), (2, "Pending", ""))
        self.assertEqual([rep.flag(0), rep.flag(1)], ["", "bad"])

    def test_adjusted_period_shows_both_subjects_and_the_original_faculty(self):
        self.set_today(week(0))
        self.adjusted(week(0), taught=self.db, topic="Joins")
        rep = builders.daily(F(date=week(0)))
        p1 = rep.rows[0]
        self.assertEqual(p1[3], "Databases (timetable: Data Structures)")
        self.assertEqual(p1[4], "Sneha (adjustment for Ravi)")
        self.assertEqual(p1[5], "Entered (adjusted)")

    def test_same_subject_adjustment_shows_one_subject(self):
        self.set_today(week(0))
        self.adjusted(week(0), same=True)
        p1 = builders.daily(F(date=week(0))).rows[0]
        self.assertEqual(p1[3], "Data Structures")
        self.assertIn("adjustment for Ravi", p1[4])

    def test_pending_adjusted_and_unadjusted_leave_periods(self):
        self.set_today(week(0))
        self.adjusted(week(0), taught=self.db)                                      # P1 adjusted; delete attendance
        from attendance.models import AttendanceSession
        AttendanceSession.objects.all().delete()
        p1 = builders.daily(F(date=week(0))).rows[0]
        self.assertEqual(p1[5], "Pending (adjusted)")
        Adjustment.objects.update(substitute=None, subject_taught=None, status=Adjustment.UNASSIGNED)
        p1 = builders.daily(F(date=week(0))).rows[0]
        self.assertEqual(p1[5], "On leave - no substitute")
        self.assertIn("Ravi (on leave)", p1[4])

    def test_holiday_is_noted_instead_of_listing_periods(self):
        Holiday.objects.create(from_date=week(0), to_date=week(0), reason="Festival")
        rep = builders.daily(F(date=week(0)))
        self.assertEqual(rep.rows, [])
        self.assertIn("A: holiday (Festival)", rep.notes)

    def test_filters(self):
        self.held(week(0), user=self.ravi)
        self.assertEqual([r[1] for r in builders.daily(F(date=week(0), faculty=self.sneha)).rows], [2])
        self.assertEqual([r[1] for r in builders.daily(F(date=week(0), subject_code="DS")).rows], [1])


class LogTests(ReportTestCase):
    def test_faculty_log_sorted_by_faculty_subject_date(self):
        self.held(week(1), topic="Queues", user=self.ravi)
        self.held(week(0), topic="Stacks", user=self.ravi)
        self.held(week(0), period=2, topic="Joins", user=self.sneha)
        self.adjusted(week(2), taught=self.db, topic="Cover for Ravi")
        rep = builders.faculty_log(F())
        self.assertEqual([(r[0], r[3], r[5]) for r in rep.rows], [
            ("Ravi", "05-10-2026", "Stacks"),
            ("Ravi", "12-10-2026", "Queues"), ("Sneha", "05-10-2026", "Joins"), ("Sneha", "19-10-2026", "Cover for Ravi")])
        cover = next(r for r in rep.rows if r[5] == "Cover for Ravi")
        self.assertEqual(cover[0], "Sneha")                       # credited to the substitute
        self.assertEqual(cover[1], "Databases (DB)")              # under the subject actually taught
        self.assertIn("Adjustment for Ravi", cover[7])
        self.assertIn("timetable subject: Data Structures", cover[7])

    def test_filters(self):
        self.held(week(0), user=self.ravi)
        self.held(week(5), user=self.ravi)
        self.assertEqual(len(builders.faculty_log(F(date_from=week(3))).rows), 1)
        self.assertEqual(len(builders.faculty_log(F(faculty=self.sneha)).rows), 0)


class LeaveRegisterTests(ReportTestCase):
    def test_register_rows_show_who_took_what_and_who_arranged_it(self):
        self.adjusted(week(0), taught=self.db)
        rep = builders.leave_register(F())
        row = rep.rows[0]
        self.assertEqual(row[:3], ["Ravi", "Casual", "Approved"])
        self.assertEqual(row[4:], [1, "A", "Data Structures", "Sneha", "Databases", "Leave cover", "Different",
                                   "Accepted", "Admin"])

    def test_same_subject_and_faculty_arranged(self):
        self.set_today(week(0) - datetime.timedelta(days=1))
        req = leave_svc.create_leave_request(faculty=self.ravi, leave_type=self.lt, from_date=week(0), to_date=week(0),
                                             scope="periods", periods="1", created_by=self.ravi)
        a = leave_svc.leave_adjustments(req).get()
        leave_svc.set_adjustment(a, user=self.ravi, substitute=self.sneha, same_subject=True)
        leave_svc.submit_request(req, user=self.ravi)
        row = builders.leave_register(F()).rows[0]
        self.assertEqual((row[2], row[9:]), ("Waiting for substitutes", ["Leave cover", "Same", "Waiting for acceptance", "Faculty"]))

    def test_leave_with_no_affected_periods_is_still_listed(self):
        self.set_today(week(0) - datetime.timedelta(days=2))
        leave_svc.create_leave_request(faculty=self.sneha, leave_type=self.lt, from_date=week(1) + datetime.timedelta(days=1),
                                       to_date=week(1) + datetime.timedelta(days=1), scope="full", created_by=self.admin)
        rows = builders.leave_register(F()).rows
        self.assertEqual([r[0] for r in rows], ["Sneha", ])

    def test_rejected_and_cancelled_leave_is_not_listed(self):
        self.adjusted(week(0), req_status=LeaveRequest.REJECTED)
        self.assertEqual(builders.leave_register(F()).rows, [])

    def test_filters(self):
        self.adjusted(week(0), taught=self.db)
        self.adjusted(week(5), taught=self.db)
        self.assertEqual(len(builders.leave_register(F(date_from=week(3))).rows), 1)
        self.assertEqual(len(builders.leave_register(F(faculty=self.sneha)).rows), 0)       # filter = faculty on leave
        self.assertEqual(len(builders.leave_register(F(faculty=self.ravi)).rows), 2)
        self.assertEqual(len(builders.leave_register(F(subject_code="DB")).rows), 2)

    def test_return_adjustment_is_marked_as_a_return(self):
        self.adjusted(week(0), taught=self.db)
        req = LeaveRequest.objects.get()
        Adjustment.objects.create(leave_request=req, kind=Adjustment.RETURN, school_class=self.c, date=week(1),
                                  period_no=2, original_faculty=self.sneha, original_subject=self.db,
                                  substitute=self.ravi, subject_taught=self.db, status=Adjustment.ACCEPTED,
                                  made_by=self.ravi, made_by_role="faculty")
        kinds = [r[9] for r in builders.leave_register(F()).rows]
        self.assertEqual(sorted(kinds), ["Leave cover", "Return (applicant takes substitute's period)"])


class WorkloadTests(ReportTestCase):
    def test_own_taken_for_others_and_given_away(self):
        self.held(week(0), user=self.ravi)                         # Ravi own
        self.held(week(0), period=2, user=self.sneha)              # Sneha own
        self.adjusted(week(1), substitute=self.sneha)              # Ravi gives away, Sneha takes
        self.adjusted(week(2), substitute=self.sneha)
        rows = {r[0]: r[1:] for r in builders.workload(F()).rows}
        self.assertEqual(rows["Ravi"], [1, 0, 2, 1])
        self.assertEqual(rows["Sneha"], [1, 2, 0, 3])
        self.assertEqual(rows["TOTAL"], [2, 2, 2, 4])

    def test_faculty_and_date_filters(self):
        self.held(week(0), user=self.ravi)
        self.adjusted(week(5), substitute=self.sneha)
        self.assertEqual([r[0] for r in builders.workload(F(faculty=self.sneha)).rows], ["Sneha", "TOTAL"])
        rows = {r[0]: r[1:] for r in builders.workload(F(date_to=week(2))).rows}
        self.assertEqual(rows["Sneha"], [0, 0, 0, 0])


class ScreenAndExportTests(ReportTestCase):
    def setUp(self):
        super().setUp()
        self.held(week(0), absent=[self.r(1)], topic="Stacks", user=self.ravi)
        self.held(week(1), absent=[self.r(1)], topic="Queues", user=self.ravi)
        self.client.force_login(self.admin)

    def test_every_report_page_renders(self):
        for rdef in builders.REPORTS:
            params = {"school_class": self.c.pk} if rdef.needs_class else {}
            r = self.client.get(reverse("report_view", args=[rdef.slug]), params)
            self.assertEqual(r.status_code, 200, rdef.slug)
        self.assertContains(self.client.get(reverse("report_index")), "Shortage list")

    def test_default_class_is_chosen_for_student_wise(self):
        self.assertContains(self.client.get(reverse("report_view", args=["student-wise"])), "Student-wise attendance")

    def test_shortage_screen_shows_the_student_and_threshold(self):
        r = self.client.get(reverse("report_view", args=["shortage"]), {"threshold": 75})
        self.assertContains(r, "S1")
        self.assertContains(r, "below 75%")
        self.assertContains(r, "0.0%")

    def test_excel_export_matches_the_screen(self):
        r = self.client.get(reverse("report_view", args=["shortage"]), {"export": "xlsx"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("shortage_", r["Content-Disposition"])
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        cells = [[c.value for c in row] for row in ws.iter_rows()]
        self.assertEqual(cells[0][0], "Attendance shortage list (below 75%)")
        header = next(i for i, row in enumerate(cells) if row[:2] == ["Class", "Roll no"])
        self.assertEqual(cells[header + 1][:6], ["A", "R1", "S1", 2, 0, "0.0%"])

    def test_pdf_export(self):
        r = self.client.get(reverse("report_view", args=["faculty-log"]), {"export": "pdf"})
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_every_report_exports_to_both_formats(self):
        for rdef in builders.REPORTS:
            for fmt in ("xlsx", "pdf"):
                params = {"export": fmt}
                if rdef.needs_class:
                    params["school_class"] = self.c.pk
                r = self.client.get(reverse("report_view", args=[rdef.slug]), params)
                self.assertEqual(r.status_code, 200, (rdef.slug, fmt))
                self.assertGreater(len(r.content), 500)

    def test_export_applies_the_filters(self):
        r = self.client.get(reverse("report_view", args=["faculty-log"]),
                            {"export": "xlsx", "date_from": week(1).isoformat()})
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        topics = [row[5].value for row in ws.iter_rows() if row[0].value == "Ravi"]
        self.assertEqual(topics, ["Queues"])

    def test_empty_report_exports_do_not_crash(self):
        from attendance.models import AttendanceSession
        AttendanceSession.objects.all().delete()
        for fmt in ("xlsx", "pdf"):
            r = self.client.get(reverse("report_view", args=["shortage"]), {"export": fmt})
            self.assertEqual(r.status_code, 200)

    def test_bad_date_range_is_reported(self):
        r = self.client.get(reverse("report_view", args=["faculty-log"]),
                            {"date_from": "2026-10-10", "date_to": "2026-10-01"})
        self.assertContains(r, "cannot be before")

    def test_unknown_report_is_404(self):
        self.assertEqual(self.client.get("/reports/nope/").status_code, 404)

    def test_faculty_cannot_open_reports_or_exports(self):
        self.client.force_login(self.ravi)
        for url in (reverse("report_index"), reverse("report_view", args=["shortage"]),
                    reverse("report_view", args=["shortage"]) + "?export=xlsx"):
            self.assertEqual(self.client.get(url).status_code, 403, url)

    def test_anonymous_is_sent_to_login(self):
        self.client.logout()
        r = self.client.get(reverse("report_view", args=["shortage"]))
        self.assertEqual(r.status_code, 302)
        self.assertIn("/login/", r["Location"])

    def test_special_characters_do_not_break_pdf(self):
        self.held(week(2), topic="Trees & <graphs> 100%", user=self.ravi)
        r = self.client.get(reverse("report_view", args=["faculty-log"]), {"export": "pdf"})
        self.assertEqual(r.status_code, 200)

    def test_threshold_setting_page(self):
        from academics.models import shortage_threshold
        self.client.post(reverse("period_settings"), {"forenoon_last_period": 4, "shortage_threshold": 80})
        self.assertEqual(shortage_threshold(), 80)
        self.client.post(reverse("period_settings"), {"forenoon_last_period": 4, "shortage_threshold": 150})
        self.assertEqual(shortage_threshold(), 80)                  # rejected
        self.client.post(reverse("period_settings"), {"forenoon_last_period": 4})
        self.assertEqual(shortage_threshold(), 80)                  # blank keeps the old value
