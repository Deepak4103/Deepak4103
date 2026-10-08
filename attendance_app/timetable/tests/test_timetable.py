import datetime
import io

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from academics.models import Allotment, SchoolClass, Student, Subject
from accounts.models import User
from attendance.tests.base import MON, AttendanceTestCase
from timetable import services as svc
from timetable.models import Timetable, TimetableEntry

HEADER = "day,period,start_time,end_time,subject,faculty\n"


def upload(client, cls, text, eff="2026-11-02", name="t.csv"):
    return client.post(reverse("timetable_upload", args=[cls.pk]),
                       {"effective_from": eff, "file": SimpleUploadedFile(name, text.encode())})


class UploadTests(AttendanceTestCase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.admin)

    def test_valid_upload_preview_then_save(self):
        r = upload(self.client, self.c, HEADER + "Mon,1,09:00,09:50,DS,ravi\nTue,1,09:00,09:50,DB,sneha\n")
        self.assertContains(r, "2 periods read")
        self.assertContains(r, "Save this timetable")
        self.assertEqual(Timetable.objects.filter(effective_from="2026-11-02").count(), 0)   # preview saves nothing
        r = self.client.post(reverse("timetable_upload", args=[self.c.pk]), {"action": "confirm"})
        tt = Timetable.objects.get(effective_from="2026-11-02")
        self.assertRedirects(r, reverse("timetable_grid", args=[tt.pk]))
        self.assertEqual(tt.entries.count(), 2)

    def test_excel_upload_with_time_cells(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["day", "period", "start_time", "end_time", "subject", "faculty"])
        ws.append(["Monday", 1, datetime.time(9, 0), datetime.time(9, 50), "Data Structures", "Ravi"])
        buf = io.BytesIO()
        wb.save(buf)
        r = self.client.post(reverse("timetable_upload", args=[self.c.pk]), {
            "effective_from": "2026-11-02", "file": SimpleUploadedFile("t.xlsx", buf.getvalue())})
        self.assertContains(r, "1 periods read")
        self.assertContains(r, "Save this timetable")

    def test_errors_block_saving(self):
        bad = HEADER + ("Mon,1,09:00,09:50,NOPE,ravi\n"       # unknown subject
                        "Fun,2,09:00,09:50,DS,ravi\n"          # unknown day
                        "Mon,9,09:00,09:50,DS,ravi\n"          # period out of range
                        "Mon,3,9am,09:50,DS,ravi\n"            # bad time
                        "Mon,4,09:00,09:50,DS,ghost\n")        # unknown faculty
        r = upload(self.client, self.c, bad)
        for msg in ("not a subject of this class", "Unknown day", "Period must be", "Cannot read time", "not found"):
            self.assertContains(r, msg)
        self.assertNotContains(r, "Save this timetable")
        r = self.client.post(reverse("timetable_upload", args=[self.c.pk]), {"action": "confirm"})
        self.assertEqual(Timetable.objects.count(), 1)

    def test_duplicate_slot_and_inconsistent_times(self):
        r = upload(self.client, self.c, HEADER + "Mon,1,09:00,09:50,DS,ravi\nMon,1,09:00,09:50,DB,sneha\n"
                                                 "Tue,2,09:50,10:40,DS,ravi\nWed,2,10:00,10:40,DS,ravi\n")
        self.assertContains(r, "listed more than once")
        self.assertContains(r, "different timings")

    def test_end_before_start(self):
        r = upload(self.client, self.c, HEADER + "Mon,1,10:00,09:00,DS,ravi\n")
        self.assertContains(r, "end time must be after start time")

    def test_faculty_clash_with_other_class(self):
        other = SchoolClass.objects.create(name="B", branch="CSE", year=2, semester=1, section="B")
        sub = Subject.objects.create(school_class=other, name="DS", code="DS")
        tt = Timetable.objects.create(school_class=other, effective_from=datetime.date(2026, 6, 1))
        TimetableEntry.objects.create(timetable=tt, day=0, period_no=3, start_time=datetime.time(11),
                                      end_time=datetime.time(12), subject=sub, faculty=self.ravi)
        r = upload(self.client, self.c, HEADER + "Mon,3,11:00,12:00,DS,ravi\n")
        self.assertContains(r, "already teaching B on Monday period 3")
        self.assertNotContains(r, "Save this timetable")

    def test_existing_version_date_rejected(self):
        r = upload(self.client, self.c, HEADER + "Mon,1,09:00,09:50,DS,ravi\n", eff=MON.isoformat())
        self.assertContains(r, "already exists from this date")

    def test_missing_header(self):
        r = upload(self.client, self.c, "a,b\n1,2\n")
        self.assertContains(r, "Header row must be")

    def test_blank_templates_download(self):
        r = self.client.get(reverse("timetable_template", args=["csv"]))
        self.assertEqual(r.content.decode().strip(), "day,period,start_time,end_time,subject,faculty")
        r = self.client.get(reverse("timetable_template", args=["xlsx"]))
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[1]], svc.HEADERS)
        self.assertEqual(ws.max_row, 1)


class ParseTests(AttendanceTestCase):
    def test_parse_helpers(self):
        self.assertEqual(svc.parse_day("Monday"), 0)
        self.assertEqual(svc.parse_day("sat"), 5)
        self.assertEqual(svc.parse_day("3"), 2)
        with self.assertRaises(ValueError):
            svc.parse_day("Sun")
        self.assertEqual(svc.parse_time("9:05"), datetime.time(9, 5))
        self.assertEqual(svc.parse_time("09:00:00"), datetime.time(9, 0))
        self.assertEqual(svc.parse_time("2:30 PM"), datetime.time(14, 30))


class GridTests(AttendanceTestCase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.admin)
        self.url = reverse("timetable_grid", args=[self.tt.pk])

    def post(self, **extra):
        data = {"effective_from": MON.isoformat()}
        for p, (s, e) in enumerate([("09:00", "10:00"), ("10:00", "11:00")], start=1):
            data[f"start_{p}"], data[f"end_{p}"] = s, e
        data.update(extra)
        return self.client.post(self.url, data)

    def test_grid_page_renders_existing_entries(self):
        r = self.client.get(self.url)
        self.assertContains(r, "Weekly grid")
        self.assertContains(r, "selected")

    def test_save_grid_and_auto_faculty_from_allotment(self):
        r = self.post(s_0_1=str(self.ds.pk), s_1_1=str(self.db.pk))       # no faculty chosen
        self.assertRedirects(r, reverse("timetable_list", args=[self.c.pk]))
        entries = {(e.day, e.period_no): e for e in self.tt.entries.all()}
        self.assertEqual(set(entries), {(0, 1), (1, 1)})
        self.assertEqual(entries[(0, 1)].faculty, self.ravi)
        self.assertEqual(entries[(1, 1)].faculty, self.sneha)
        self.assertEqual(entries[(0, 1)].end_time, datetime.time(10, 0))

    def test_clearing_a_slot_removes_the_entry(self):
        self.post(s_0_1=str(self.ds.pk))
        self.assertEqual(self.tt.entries.count(), 1)

    def test_invalid_times_rejected(self):
        r = self.post(s_0_1=str(self.ds.pk), start_1="", end_1="")
        self.assertContains(r, "valid start and end times")
        self.assertEqual(self.tt.entries.count(), 2)   # unchanged

    def test_clash_with_other_class_rejected(self):
        other = SchoolClass.objects.create(name="B", branch="CSE", year=2, semester=1, section="B")
        sub = Subject.objects.create(school_class=other, name="DS", code="DS")
        tt = Timetable.objects.create(school_class=other, effective_from=MON)
        TimetableEntry.objects.create(timetable=tt, day=1, period_no=1, start_time=datetime.time(9),
                                      end_time=datetime.time(10), subject=sub, faculty=self.ravi)
        r = self.post(s_1_1=str(self.ds.pk), f_1_1=str(self.ravi.pk))
        self.assertContains(r, "already teaching B")

    def test_duplicate_version_date_rejected(self):
        other = Timetable.objects.create(school_class=self.c, effective_from=datetime.date(2026, 12, 1))
        r = self.post(effective_from="2026-12-01", s_0_1=str(self.ds.pk))
        self.assertContains(r, "Another version already starts on that date")

    def test_new_version_copy_and_delete(self):
        r = self.client.post(reverse("timetable_new", args=[self.c.pk]),
                             {"effective_from": "2026-12-01", "copy_from": self.tt.pk})
        new = Timetable.objects.get(effective_from="2026-12-01")
        self.assertEqual(new.entries.count(), 2)
        self.client.post(reverse("timetable_delete", args=[new.pk]))
        self.assertFalse(Timetable.objects.filter(pk=new.pk).exists())

    def test_cannot_delete_subject_used_in_timetable(self):
        r = self.client.post(reverse("subject_delete", args=[self.ds.pk]), follow=True)
        self.assertContains(r, "Cannot delete")
        self.assertTrue(Subject.objects.filter(pk=self.ds.pk).exists())
